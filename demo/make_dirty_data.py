#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成演示用脏表格。每处"脏"都对应 references/excel-pitfalls.md 里的一个坑。

用法:
    python make_dirty_data.py [输出路径]
"""
from __future__ import annotations

import os
import random
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

REGIONS = ["华东", "华南", "华北", "西南"]
CHANNELS = ["直营", "经销", "电商"]
PRODUCTS = ["A100", "B200", "C300"]


def build(out_path: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "销售明细"

    # 坑1: 前置装饰行——真表头不在第 1 行
    ws["A1"] = "2026年 区域销售明细表（内部资料）"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:G1")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "制表人：尚威"
    ws["D2"] = "更新时间：2026/09/16"
    # 第 3 行故意留空

    header = ["订单日期", "区域", "渠道", "产品型号", "销售数量", "销售额", "备注"]
    for col, name in enumerate(header, start=1):
        c = ws.cell(row=4, column=col, value=name)
        c.font = Font(bold=True)

    random.seed(42)
    rows: list[list] = []
    for i in range(60):
        month = random.randint(1, 9)
        day = random.randint(1, 28)
        qty = random.randint(1, 50)
        amount = qty * random.choice([1280, 2350, 4600])

        # 坑2: 日期以文本存储，且格式不统一
        if i % 3 == 0:
            date = f"2026/{month}/{day}"
        elif i % 3 == 1:
            date = f"2026-{month:02d}-{day:02d}"
        else:
            date = f"2026年{month}月{day}日"

        region = random.choice(REGIONS)
        # 坑3: 首尾空格 + 全角空格污染，导致分组时同名被拆成多组
        if i % 7 == 0:
            region = f" {region} "
        elif i % 11 == 0:
            region = f"{region}\u3000"

        # 坑4: 金额带千分位和货币符号，被 Excel 当文本
        amount_cell = f"¥{amount:,}" if i % 4 else amount

        # 坑5: 数量偶尔是文本
        qty_cell = str(qty) if i % 9 == 0 else qty

        rows.append([date, region, random.choice(CHANNELS),
                     random.choice(PRODUCTS), qty_cell, amount_cell,
                     None if i % 5 else "加急"])

    # 坑6: 重复行
    rows.insert(20, list(rows[8]))
    rows.insert(35, list(rows[8]))

    # 坑7: 中间夹空行
    rows.insert(15, [None] * 7)
    rows.insert(42, [None] * 7)

    for r_idx, row in enumerate(rows, start=5):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    # 坑8: 表格末尾的合计行——会被误当成数据行参与统计
    last = 5 + len(rows) + 1
    ws.cell(row=last, column=1, value="合计")
    ws.cell(row=last, column=5, value=f"=SUM(E5:E{5 + len(rows) - 1})")
    ws.cell(row=last, column=1).font = Font(bold=True)

    # 坑9: 尾部整列为空的幽灵列
    ws.cell(row=4, column=9, value=None)
    ws.cell(row=6, column=9, value=" ")

    # 坑10: 第二个 sheet 是无关的说明页，盲目取 sheet 会取错
    ws2 = wb.create_sheet("填表说明")
    ws2["A1"] = "1. 销售额含税"
    ws2["A2"] = "2. 区域按大区划分"

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path


def main() -> int:
    out = sys.argv[1] if len(sys.argv) > 1 else \
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "dirty_sales.xlsx")
    path = build(out)
    print(f"已生成演示脏表: {path}")
    print("包含 10 类典型脏数据：装饰行/文本日期/空格污染/千分位金额/"
          "文本数字/重复行/空行/合计行/幽灵列/多 sheet 干扰")
    return 0


if __name__ == "__main__":
    sys.exit(main())
